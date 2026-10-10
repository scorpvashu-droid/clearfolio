"""Local, rule-based resume text extraction and parsing.

The parser keeps the original project's Resume/Experience data shape and
section-aware extraction approach without requiring an external AI service.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable
from pathlib import Path

import pymupdf
from docx import Document
from pydantic import BaseModel, ConfigDict, Field


class Education(BaseModel):
    institution: str | None = None
    degree: str | None = None
    field: str | None = None
    dates: str | None = None
    details: str | None = None


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    job_title: str | None = None
    duration: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Skills(BaseModel):
    technical: list[str] = Field(default_factory=list)
    soft: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str | None = None
    email: str | None = None
    phone: str | None = None
    location: str | None = None
    summary: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = Field(default_factory=list)
    technical_skills: list[str] = Field(default_factory=list)
    soft_skills: list[str] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    languages: list[str] = Field(default_factory=list)


class ResumeTextError(ValueError):
    """Raised when a document contains no usable text."""


class OCRRequiredError(ResumeTextError):
    """Raised when a PDF contains no embedded text and requires OCR."""


_SECTION_ALIASES: dict[str, str] = {
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "professional profile": "summary",
    "objective": "summary",
    "career objective": "summary",
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "employment history": "experience",
    "work history": "experience",
    "internships": "experience",
    "education": "education",
    "academic background": "education",
    "skills": "skills",
    "technical skills": "skills",
    "core competencies": "skills",
    "competencies": "skills",
    "projects": "projects",
    "personal projects": "projects",
    "certifications": "certifications",
    "certificates": "certifications",
    "licenses and certifications": "certifications",
    "languages": "languages",
}

_TECHNICAL_SKILLS = (
    "Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "Go", "Rust",
    "SQL", "HTML", "CSS", "React", "Angular", "Vue", "Node.js", "FastAPI",
    "Django", "Flask", "Spring", "AWS", "Azure", "GCP", "Docker", "Kubernetes",
    "Git", "Linux", "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis",
    "TensorFlow", "PyTorch", "Pandas", "NumPy", "scikit-learn", "REST",
    "GraphQL", "CI/CD", "Jenkins", "Terraform", "Excel", "Power BI", "Tableau",
)
_SOFT_SKILLS = (
    "communication", "leadership", "teamwork", "problem solving",
    "problem-solving", "adaptability", "collaboration", "time management",
    "critical thinking", "creativity", "mentoring", "organization",
    "attention to detail", "conflict resolution", "customer service",
)
_DATE = re.compile(
    r"\b(?P<start>(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?\.?\s+)?(?:19|20)\d{2})\s*"
    r"(?:[-–—]|to)\s*(?P<end>(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|"
    r"Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|"
    r"Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?\.?\s+)?(?:19|20)\d{2}|"
    r"Present|Current|Now)\b",
    re.IGNORECASE,
)
_EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
_PHONE = re.compile(r"(?<!\w)(?:\+?\d[\d().\-\s]{6,}\d)(?!\w)")
_DEGREE = re.compile(
    r"\b(Ph\.?D\.?|Doctor(?:ate)?|M\.?S\.?|M\.?A\.?|M\.?B\.A\.?|"
    r"B\.?S\.?|B\.?A\.?|B\.?E\.?|B\.?Tech|Bachelor(?:'s)?|Master(?:'s)?|"
    r"Associate(?:'s)?|Diploma)\b",
    re.IGNORECASE,
)
_TITLE = re.compile(
    r"\b(engineer|developer|manager|analyst|designer|consultant|intern|"
    r"architect|administrator|specialist|director|scientist|lead|coordinator|"
    r"officer|technician|researcher|product owner|accountant)\b",
    re.IGNORECASE,
)


def read_pdf(source: bytes | bytearray | io.BytesIO | str | Path) -> str:
    """Extract embedded text from a PDF; scanned PDFs are not OCR'd."""
    if isinstance(source, (str, Path)):
        document = pymupdf.open(source)
    else:
        raw = source.getvalue() if isinstance(source, io.BytesIO) else bytes(source)
        document = pymupdf.open(stream=raw, filetype="pdf")
    with document:
        text = "\n".join(page.get_text("text") for page in document)
    if not text.strip():
        raise OCRRequiredError(
            "This PDF has no extractable text. It may be scanned; OCR is required."
        )
    return text


def read_docx(source: bytes | bytearray | io.BytesIO | str | Path) -> str:
    """Extract paragraph and table text from a DOCX document."""
    document = Document(source if isinstance(source, (str, Path)) else io.BytesIO(bytes(source)))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    return "\n".join(parts)


def read_resume(source: bytes | bytearray | io.BytesIO | str | Path, filename: str | None = None) -> str:
    """Read a PDF or DOCX from bytes or a path, selected by its filename."""
    suffix = (filename or str(source)).lower()
    if suffix.endswith(".pdf"):
        return read_pdf(source)
    if suffix.endswith(".docx"):
        return read_docx(source)
    raise ValueError("Only PDF and DOCX resumes are supported.")


def _clean_line(line: str) -> str:
    return re.sub(r"^\s*(?:[•*▪◦\-–—]|\d+[.)])\s*", "", line).strip()


def _section_key(line: str) -> str | None:
    normalized = re.sub(r"[:\s]+$", "", line.strip().lower())
    normalized = re.sub(r"^\d+[.)]\s*", "", normalized)
    return _SECTION_ALIASES.get(normalized)


def _sections(text: str) -> tuple[list[str], dict[str, list[str]]]:
    top_lines: list[str] = []
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in text.splitlines():
        cleaned = _clean_line(line)
        section = _section_key(cleaned)
        if section is not None:
            current = section
            sections.setdefault(current, [])
        elif current is not None:
            sections[current].append(cleaned)
        elif cleaned:
            top_lines.append(cleaned)
    return top_lines, sections


def _section_text(section: Iterable[str]) -> str | None:
    value = " ".join(line for line in section if line).strip()
    return value or None


def _unique(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = value.strip(" \t,;|")
        key = cleaned.casefold()
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return result


def _items(lines: Iterable[str]) -> list[str]:
    values: list[str] = []
    for line in lines:
        cleaned = _clean_line(line)
        if cleaned:
            values.extend(re.split(r"\s*[;|]\s*|\s*,\s*", cleaned))
    return _unique(values)


def _recognized_skills(text: str, vocabulary: Iterable[str]) -> list[str]:
    return [
        skill for skill in vocabulary
        if re.search(rf"(?<![\w+#.]){re.escape(skill)}(?![\w+#.])", text, re.IGNORECASE)
    ]


def _contact_location(lines: list[str]) -> str | None:
    for line in lines[:8]:
        match = re.search(r"\b(?:location|based in|address)\s*[:\-]\s*(.+)$", line, re.IGNORECASE)
        if match:
            return match.group(1).strip(" ,|")
    for line in lines[:5]:
        if not (_EMAIL.search(line) or _PHONE.search(line)):
            continue
        for part in re.split(r"\s*[|•]\s*", line):
            candidate = _EMAIL.sub("", _PHONE.sub("", part)).strip(" ,|")
            if "," in candidate and not re.search(r"\d", candidate):
                return candidate
    return None


def _chunks(lines: list[str]) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if not line:
            if current:
                chunks.append(current)
                current = []
            continue
        has_date = any(_DATE.search(existing) for existing in current)
        begins_new_entry = (
            has_date
            and len(line) < 100
            and not re.search(r"[.!?]$", line)
            and (_TITLE.search(line) or _DEGREE.search(line))
        )
        if current and (begins_new_entry or (_DATE.search(line) and has_date)):
            chunks.append(current)
            current = []
        current.append(line)
    if current:
        chunks.append(current)
    return chunks


def _parse_experience(lines: list[str]) -> list[Experience]:
    results: list[Experience] = []
    for chunk in _chunks(lines):
        raw = " ".join(chunk)
        date_match = _DATE.search(raw)
        labels = {
            key: re.search(rf"\b{key}\s*:\s*(.+)", raw, re.IGNORECASE)
            for key in ("company", "employer", "title", "role", "position")
        }
        company = next(
            (match.group(1).split("|")[0].strip() for key, match in labels.items()
             if key in {"company", "employer"} and match),
            None,
        )
        title = next(
            (match.group(1).split("|")[0].strip() for key, match in labels.items()
             if key in {"title", "role", "position"} and match),
            None,
        )
        metadata_lines: set[str] = set()
        for line in chunk:
            clean_line = _DATE.sub("", line).strip(" -–—")
            relationship = re.match(r"(?P<title>.+?)\s+(?:at|@)\s+(?P<company>.+)", clean_line, re.I)
            if relationship:
                title = title or relationship.group("title").strip()
                company = company or relationship.group("company").strip()
                metadata_lines.add(line)
                continue
            if title is None and _TITLE.search(clean_line):
                title = clean_line
                metadata_lines.add(line)
            if company is None and re.search(
                r"\b(?:Inc\.?|LLC|Ltd\.?|Corporation|Corp\.?|Technologies|"
                r"University|Company|Group|Solutions|Systems)\b", clean_line, re.IGNORECASE
            ):
                company = clean_line
                metadata_lines.add(line)
        metadata_lines.update(
            line for line in chunk
            if re.match(r"^\s*(?:company|employer|title|role|position)\s*:", line, re.IGNORECASE)
        )
        description_lines = [
            line for line in chunk
            if line not in metadata_lines and not _DATE.search(line)
        ]
        description = " ".join(description_lines).strip() or None
        results.append(Experience(
            company=company,
            role=title,
            job_title=title,
            duration=date_match.group(0) if date_match else None,
            start_date=date_match.group("start") if date_match else None,
            end_date=date_match.group("end") if date_match else None,
            description=description,
            skills_used=_recognized_skills(raw, _TECHNICAL_SKILLS),
        ))
    return results


def _parse_education(lines: list[str]) -> list[Education]:
    result: list[Education] = []
    for chunk in _chunks(lines):
        raw = " ".join(chunk)
        degree_match = _DEGREE.search(raw)
        degree = degree_match.group(0) if degree_match else None
        institution = next(
            (line for line in chunk if re.search(
                r"\b(university|college|institute|school|academy)\b", line, re.IGNORECASE
            )),
            None,
        )
        field = None
        if degree_match:
            field_match = re.search(r"\b(?:in|of)\s+([A-Z][A-Za-z &/-]+)", raw[degree_match.end():])
            field = field_match.group(1).strip() if field_match else None
        date_match = _DATE.search(raw)
        result.append(Education(
            institution=institution,
            degree=degree,
            field=field,
            dates=date_match.group(0) if date_match else None,
            details=raw or None,
        ))
    return result


def parse_resume(resume_text: str) -> Resume:
    """Parse resume text conservatively; absent or uncertain fields remain empty."""
    if not resume_text or not resume_text.strip():
        raise ResumeTextError("No readable text was found in the resume.")

    top_lines, sections = _sections(resume_text)
    candidate_name = next(
        (line for line in top_lines[:6]
         if not _EMAIL.search(line) and not _PHONE.search(line)
         and not re.search(r"\b(resume|curriculum vitae|linkedin\.com|github\.com)\b", line, re.I)
         and not re.match(r"\s*(?:location|based in|address)\s*[:\-]", line, re.I)
         and len(line) <= 70 and not re.search(r"[@:/]", line)),
        None,
    )

    explicit_skills = _items(sections.get("skills", []))
    soft = _recognized_skills(resume_text, _SOFT_SKILLS)
    technical = _unique(
        explicit_skills + _recognized_skills(resume_text, _TECHNICAL_SKILLS)
    )
    skills = _unique(technical + soft)
    email = next(iter(_EMAIL.findall(resume_text)), None)
    phone_match = _PHONE.search(resume_text)
    phone = re.sub(r"\s+", " ", phone_match.group(0)).strip() if phone_match else None

    explicit_experience = re.search(
        r"\b(\d+(?:\.\d+)?)\+?\s+years?\s+(?:of\s+)?experience\b",
        resume_text,
        re.IGNORECASE,
    )
    total_experience = float(explicit_experience.group(1)) if explicit_experience else None

    projects = _items(sections.get("projects", []))
    certifications = _items(sections.get("certifications", []))
    languages = _items(sections.get("languages", []))
    return Resume(
        name=candidate_name,
        email=email,
        phone=phone,
        location=_contact_location(top_lines),
        summary=_section_text(sections.get("summary", [])),
        total_experience_years=total_experience,
        skills=skills,
        technical_skills=technical,
        soft_skills=soft,
        education=_parse_education(sections.get("education", [])),
        experiences=_parse_experience(sections.get("experience", [])),
        projects=projects,
        certifications=certifications,
        languages=languages,
    )


def parse_resume_file(content: bytes, filename: str) -> Resume:
    """Extract and parse a supported resume file without writing it to disk."""
    text = read_resume(content, filename)
    return parse_resume(text)
